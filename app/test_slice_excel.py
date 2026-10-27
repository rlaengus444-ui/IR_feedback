import json,tempfile,unittest,zipfile
from pathlib import Path
from unittest.mock import Mock
from openpyxl import load_workbook
from PIL import Image
from slice_data import SliceConfig,analyze,FEATURE_NAMES
from slice_workbook import save_workbook,TEMPLATE
from excel_workbook import WorkbookSpec,inspect_workbook
import test_excel_workbook as excel_test_helpers
from test_excel_workbook import workbook,spec

class SliceExcelTests(unittest.TestCase):
    def test_new_fields_are_distinct_same_row_and_cannot_overwrite_formula(self):
        for start in ['C13','D13','BL14','XFD13']:
            with self.assertRaises(ValueError):spec(slice_start_cell=start).validate()
        with self.assertRaises(ValueError):inspect_workbook(spec(slice_start_cell='D13'),workbook())
    def test_template_default_control_formula_preserved_and_extra_inputs_declared(self):
        s=WorkbookSpec(str(TEMPLATE),'수동 테스트','B13','C13','D13','B5','A13','BL13')
        m=inspect_workbook(s,TEMPLATE.read_bytes());self.assertEqual(3012,m['last_row'])
        book=load_workbook(TEMPLATE);self.assertIn('C13+V13',book['수동 테스트']['D13'].value)
        self.assertEqual(12,len(s.slice_cells()));self.assertEqual(('dpi_y','BW13'),s.slice_cells()[-1])
    def test_offline_export_clears_unknown_values_and_can_be_loaded_by_existing_controller(self):
        with tempfile.TemporaryDirectory() as directory:
            root=Path(directory);paths=[]
            for n in (1,3):
                p=root/f'{n:03}.png';Image.new('L',(20,20),255 if n==3 else 0).save(p);paths.append(p)
            d=analyze(paths,SliceConfig(2,25.4,25.4,first_layer=1189));out=root/'manual.xlsx';save_workbook(out,d)
            b=load_workbook(out);s=b['수동 테스트']
            self.assertEqual(1189,s['A13'].value);self.assertEqual(1,s['BL13'].value);self.assertEqual(0,s['BO13'].value)
            self.assertEqual(0,s['BL14'].value);self.assertIsNone(s['BO14'].value)
            self.assertEqual(100,s['BM15'].value);self.assertEqual(0,s['BP15'].value)
            w=WorkbookSpec(str(out),'수동 테스트','B13','C13','D13','B5','A13','BL13')
            inspect_workbook(w,out.read_bytes())
    def test_worker_payload_contains_fresh_current_and_previous_values_and_rejects_mismatch(self):
        c=excel_test_helpers.WorkbookControlTest().controller(60.)
        c.spec=spec(slice_start_cell='BL13')
        c.process(1,65.,60.)
        image=c.calls[0]['image'];self.assertEqual(0,image['image_valid']);self.assertIsNone(image['center_coverage_pct'])
        invalid=dict(image,image_valid=1)
        with self.assertRaises(ValueError):c.process(2,65.,60.,slice_features=invalid)
        values=dict(image,image_valid=1,center_coverage_pct=25.,center_intensity_pct=10.,whole_coverage_pct=50.)
        c.process(2,65.,60.,slice_features=values)
        self.assertEqual(25.,c.calls[-1]['image']['center_coverage_pct'])
    def test_image_formulas_in_input_cells_are_rejected(self):
        data=workbook(extra='<c r="E13"><f>1+1</f></c>')
        with self.assertRaises(ValueError):inspect_workbook(spec(slice_start_cell='E13'),data)
if __name__=='__main__':unittest.main()
