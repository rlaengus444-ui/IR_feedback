"""Offline staged/frozen-row target, manual setting and layer-gap parity."""
from pathlib import Path
import argparse
import json
import sys
from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

from crosscheck_image_guard import recalculate, APP
sys.path.insert(0, str(APP))
from layer_algorithm import LayerAdaptiveController, LayerSettings
from slice_data import FEATURE_NAMES
from slice_workbook import TEMPLATE


def features():
    return dict(zip(FEATURE_NAMES, (1, 10, 10, 10, 1, 10, 10, 10, 0, 5, 400, 400)))


def fill(sheet, row, layer, temperature, actual):
    for col, value in zip(('A', 'B', 'C'), (layer, temperature, actual)):
        sheet[f'{col}{row}'] = value
    for index, name in enumerate(FEATURE_NAMES, 64):
        sheet.cell(row, index, features()[name])


def compare(path, expected):
    cached = load_workbook(path, data_only=True)
    sheet = cached['수동 테스트']
    checks = 0
    for row, record in expected.items():
        for col, value in record.items():
            cell = sheet.cell(row, column_index_from_string(col))
            if value is None:
                assert cell.value in (None, ''), (row, col, value, cell.value)
            else:
                assert isinstance(cell.value, (int, float)), (row, col, value, cell.value)
                assert abs(cell.value - value) < 1e-7, (row, col, value, cell.value)
            checks += 1
    for page in cached:
        for cells in page.iter_rows():
            assert not any(c.data_type == 'e' for c in cells), (page.title, cells[0].row)
    cached.close()
    return checks


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--output', type=Path, required=True)
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=False)
    settings = LayerSettings.from_dict(json.loads((APP / 'workbook-replay.settings.json').read_text()))
    controller = LayerAdaptiveController(settings)
    book = load_workbook(TEMPLATE)
    sheet = book['수동 테스트']
    sheet.delete_rows(33, 3000)
    actual = 60
    expected = {}
    for layer, temperature in enumerate((65., 65.1, 65., 59., 59.2, 59.4), 1):
        fill(sheet, 12 + layer, layer, temperature, actual)
        result = controller.process(layer, temperature, actual, slice_features=features())
        expected[12 + layer] = result.diagnostics
        actual = result.power_percent
    source = args.output / 'initial.xlsx'
    book.save(source)
    converted = args.output / 'calc-initial'
    recalculate([source], converted)
    current = converted / source.name
    total = compare(current, expected)
    stages = []
    for index, (name, layer, temperature, target, manual_delta) in enumerate((
            ('target', 7, 59.7, 70., 0.),
            ('manual', 8, 59.8, 70., 1.),
            ('gap', 10, 59.9, 70., 0.)), 1):
        # Freeze every preceding used cell to its cached value, as excel_worker.ps1
        # does. The next row retains formulas. This is not an Excel COM test.
        formulas = load_workbook(current)
        values = load_workbook(current, data_only=True)
        dst, src = formulas['수동 테스트'], values['수동 테스트']
        for row in range(13, 18 + index):
            for cells in src.iter_rows(min_row=row, max_row=row):
                for cell in cells:
                    dst.cell(row, cell.column).value = cell.value
        row = 18 + index
        dst['B5'] = target
        controller.set_target_temperature(target)
        actual += manual_delta
        fill(dst, row, layer, temperature, actual)
        result = controller.process(layer, temperature, actual, slice_features=features())
        actual = result.power_percent
        expected[row] = result.diagnostics
        assert result.diagnostics['O'] == 1 and result.diagnostics['CI'] == 0
        source = args.output / (name + '.xlsx')
        formulas.save(source)
        values.close()
        formulas.close()
        converted = args.output / ('calc-' + name)
        recalculate([source], converted)
        current = converted / source.name
        total += compare(current, expected)
        stages.append(dict(case=name, layer=layer, row=row, power=actual))
    result = dict(stages=stages, field_checks=total, error_cells=0,
                  engine='LibreOffice staged recalculation with frozen historical rows',
                  excel_com_tested=False, equipment_tested=False)
    (args.output / 'results.json').write_text(json.dumps(result, indent=2) + '\n')
    print(json.dumps(result, indent=2))


if __name__ == '__main__':
    main()
