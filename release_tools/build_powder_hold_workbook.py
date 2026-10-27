"""Build the separate powder-control/binder-hold workbook from pinned slice.2."""
import argparse, hashlib
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

SOURCE_SHA256='7d870e35c7c288573bf5ab279ede6eaab8fe37633cd0afbcc2c389bcdd79c8b3'

def build(source, output):
    source=Path(source);output=Path(output)
    if hashlib.sha256(source.read_bytes()).hexdigest()!=SOURCE_SHA256: raise ValueError('unexpected source workbook')
    b=load_workbook(source);s=b['수동 테스트']
    labels={'CD':'현재 계획 -1미확인/0분말/1분사','CE':'직전 계획 -1/0/1','CF':'분사 기준 실제 IR %','CG':'기준 연속 확인 1/0','CH':'구간 시작 G','CI':'구간 시작 알파','CJ':'구간 시작 베타','CK':'이미지 입력 유효'}
    for col,label in labels.items():
        cell=s[f'{col}12'];cell.value=label;cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='365C7C');cell.alignment=Alignment(wrap_text=True);s.column_dimensions[col].width=22
    for r in range(13,3013):
        p=r-1;valid=f'N{r}="입력 완료"'
        parts=[]
        for flag,start,end in [('BL','BM','BO'),('BP','BQ','BS')]:
            parts.append(f'OR(AND(OR(ISBLANK({flag}{r}),{flag}{r}=0),COUNTA({start}{r}:{end}{r})=0),AND(ISNUMBER({flag}{r}),{flag}{r}=1,COUNT({start}{r}:{end}{r})=3,MIN({start}{r}:{end}{r})>=0,MAX({start}{r}:{end}{r})<=100))')
        parts.append(f'OR(ISBLANK(BT{r}),AND(ISNUMBER(BT{r}),BT{r}>=0,BT{r}=INT(BT{r})))')
        for col in ('BU','BV','BW'): parts.append(f'OR(ISBLANK({col}{r}),AND(ISNUMBER({col}{r}),{col}{r}>0))')
        s[f'CK{r}']='=AND('+','.join(parts)+')'
        old=s[f'N{r}'].value[1:];s[f'N{r}']=f'=IF(B{r}=-1,"온도계 비활성(-1)",IF(NOT(CK{r}),"이미지 입력 확인",{old}))'
        def put(col,expr): s[f'{col}{r}']=f'=IF({valid},{expr},"")'
        put('CD',f'IF(BL{r}<>1,-1,IF(BO{r}>0,1,0))')
        job='TRUE' if r==13 else f'OR(G{r}="새로 시작",N{p}<>"입력 완료",A{r}<>A{p}+1)'
        put('CE','-1' if r==13 else f'IF({job},-1,CD{p})')
        put('O',f'IF(OR({job},CD{r}<>0,CE{r}<>0),1,MIN(3,O{p}+1))' if r>13 else '1')
        for col,globalcell,lastcol in [('CH','F7','L'),('CI','B7','J'),('CJ','D7','K')]:
            globalref=f"'수동 테스트'!${globalcell[0]}${globalcell[1:]}"
            value=globalref if r==13 else f'IF({job},{globalref},IF(OR(CD{r}<>0,CE{r}<>0),{lastcol}{p},{col}{p}))'
            put(col,value)
        put('CF',f'IF(CD{r}=1,C{r},"")')
        verify='0' if r==13 else f'IF(CD{r}<>1,0,IF({job},0,IF(CE{r}=0,1,IF(CE{r}=1,IF(CG{p}=1,IF(ABS(C{r}-CF{p})<=0.000001,1,0),0),0))))'
        put('CG',verify)
        for col in range(4,64):
            cell=s.cell(r,col);value=cell.value
            if not isinstance(value,str) or not value.startswith('=') or col in (6,14,15):continue
            for ref,repl in (("'수동 테스트'!$F$7",f'CH{r}'),("'수동 테스트'!$B$7",f'CI{r}'),("'수동 테스트'!$D$7",f'CJ{r}')): value=value.replace(ref,repl)
            cell.value=value
        old_d=s[f'D{r}'].value[1:];s[f'D{r}']=f'=IF({valid},IF(CD{r}<>0,C{r},{old_d}),"")'
        old_f=s[f'F{r}'].value[1:];s[f'F{r}']=f'=IF({valid},IF(CD{r}=1,"분사층: 현재 IR 고정",IF(CD{r}=-1,"이미지 미확인: 현재 IR 고정",{old_f})),IF(N{r}="","",N{r}))'
    guide=b['이미지 입력 사용법'];guide.delete_rows(1,guide.max_row)
    for line in ('분말층 G·알파·베타 제어 / 분사층 현재 IR 고정 시험 시트',
                 '현재 이미지 전체 도포율 0은 분말층, 0 초과는 분사층입니다. 유효한 이미지가 없으면 미확인으로 보고 현재 IR을 유지합니다.',
                 '분사층 진입 시 추천값이 아닌 현재 저장된 실제 IR을 유지합니다. 수동으로 변경하면 새 실제값을 유지하고 연속 확인을 해제합니다.',
                 '현재 이미지는 현재 완료층의 분사 계획이며 온도는 그 층 분사 전 새 측정값입니다. 권장 IR은 그 층 프린팅 후 계산됩니다. 이미지는 실제 분사나 건조를 증명하지 않습니다.',
                 '기존 실행파일에서 Excel 수식 파일 → 찾아보기를 선택하고 다운로드한 이 XLSX를 직접 여세요. 시트 수동 테스트, 온도 B13, 현재 IR C13, 결과 D13, 목표 B5, 층 A13, 이미지 BL13을 지정합니다.',
                 '기존 실행파일의 수동 계산 시트 저장 버튼은 이전 참고용 서식을 저장하므로 이 파일을 열 때 사용하지 마세요.',
                 '분사층과 미확인 구간은 G·알파·베타를 학습하지 않습니다. 분말층으로 돌아오면 마지막 계수를 보존하고 세 관찰부터 새로 계산합니다.',
                 '목표보다 온도가 높아도 분사층에서는 현재 IR을 유지하는 비교용 정책입니다. 실장비 안전성은 검증하지 않았으며 수동 검토가 필요합니다.'):
        guide.append([line]);guide.cell(guide.max_row,1).alignment=Alignment(wrap_text=True);guide.row_dimensions[guide.max_row].height=40
    b.calculation.fullCalcOnLoad=True;b.save(output)
if __name__=='__main__':
    p=argparse.ArgumentParser();p.add_argument('source');p.add_argument('output');args=p.parse_args();build(args.source,args.output)
