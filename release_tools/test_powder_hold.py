import json,sys,unittest
from pathlib import Path
from dataclasses import replace
ROOT=Path(__file__).resolve().parents[1];sys.path.insert(0,str(ROOT/'app'));sys.path.insert(0,str(ROOT/'release_tools'))
from layer_algorithm import LayerSettings,LayerAdaptiveController
from powder_hold_reference import PowderHoldReference

def features(phase):
    ink = None if phase == -1 else 0 if phase == 0 else 50
    return dict(image_valid=int(ink is not None),center_coverage_pct=ink,center_intensity_pct=ink,whole_coverage_pct=ink,
                previous_image_valid=0,previous_center_coverage_pct=None,previous_center_intensity_pct=None,previous_whole_coverage_pct=None,
                previous_powder_layers=0,radius_mm=5,dpi_x=400,dpi_y=400)

def settings(): return LayerSettings.from_dict(json.loads((ROOT/'app/workbook-replay.settings.json').read_text(encoding='utf-8')))

class PowderHoldTests(unittest.TestCase):
    def test_original_80_and_67_all_fields(self):
        for name in ('random80_validation.json','input_example_validation.json'):
            fixture=json.loads((ROOT/'app'/name).read_text(encoding='utf-8'));c=PowderHoldReference(settings())
            for row in fixture['rows']:
                actual=c.process(row['A'],row['B'],row['C'],restart=row.get('reset',False),slice_features=features(0)).diagnostics
                for key,value in row.items():
                    if key=='reset':continue
                    if value is None:self.assertIsNone(actual[key],(name,row['A'],key))
                    else:self.assertAlmostEqual(value,actual[key],delta=max(1e-9,abs(value)*1e-10),msg=str((name,row['A'],key)))

    def test_hold_actual_manual_unknown_and_return(self):
        c=PowderHoldReference(settings())
        for layer in range(1,4): r=c.process(layer,50,60,slice_features=features(0))
        self.assertGreater(r.power_percent,60)
        r=c.process(4,90,60,slice_features=features(1));self.assertEqual((60,60,1),(r.power_percent,r.diagnostics['CF'],r.diagnostics['CG']))
        r=c.process(5,90,63,slice_features=features(1));self.assertEqual((63,63,0),(r.power_percent,r.diagnostics['CF'],r.diagnostics['CG']))
        r=c.process(6,90,63,slice_features=features(1));self.assertEqual((63,0),(r.power_percent,r.diagnostics['CG']))
        r=c.process(7,80,62,slice_features=features(-1));self.assertEqual((62,None,0),(r.power_percent,r.diagnostics['CF'],r.diagnostics['CG']))
        r=c.process(8,40,61,slice_features=features(1));self.assertEqual((61,0),(r.power_percent,r.diagnostics['CG']))
        c.set_target_temperature(66);r=c.process(9,95,61,slice_features=features(1));self.assertEqual(61,r.power_percent)
        for layer,phase,count in ((10,0,1),(11,0,2),(12,0,3)):
            r=c.process(layer,60,61,slice_features=features(phase));self.assertEqual(count,r.diagnostics['O'])

    def test_segment_preserves_learning_and_gap_reset(self):
        c=PowderHoldReference(settings());r=c.process(1,65,60,slice_features=features(0))
        c.controller._history[-1].update(L=.8,J=.3,K=.4)
        c.last.update(L=.8,J=.3,K=.4)
        r=c.process(2,65,60,slice_features=features(1))
        self.assertEqual((.8,.3,.4),(r.diagnostics['L'],r.diagnostics['J'],r.diagnostics['K']))
        r=c.process(3,65,60,slice_features=features(0))
        self.assertEqual((1,None,0,0),(r.diagnostics['O'],r.diagnostics['X'],r.diagnostics['AM'],r.diagnostics['AV']))
        self.assertEqual((.8,.3,.4),(r.diagnostics['L'],r.diagnostics['J'],r.diagnostics['K']))
        r=c.process(5,65,60,slice_features=features(0))
        self.assertEqual((settings().initial_g,settings().initial_alpha,settings().initial_beta),(r.diagnostics['L'],r.diagnostics['J'],r.diagnostics['K']))

    def test_invalid_no_advance(self):
        c=PowderHoldReference(settings());c.process(1,65,60,slice_features=features(0))
        for temp,f in ((-1,features(1)),(65,dict(features(1),whole_coverage_pct=float('nan')))):
            with self.assertRaises(ValueError):c.process(2,temp,60,slice_features=f)
        self.assertEqual(2,c.process(2,65,60,slice_features=features(0)).diagnostics['O'])

if __name__=='__main__' and (len(sys.argv)==1 or sys.argv[1]!='--parity'):unittest.main()

def run_parity(output):
    import csv,hashlib,shutil,subprocess,zipfile,xml.etree.ElementTree as ET
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string
    from slice_data import analyze,SliceConfig,FEATURE_NAMES
    out=Path(output);out.mkdir(parents=True,exist_ok=False)
    template=ROOT/'app/templates/IR_파우더층제어_분사층고정.xlsx'
    directory=Path('/workspace/ir-image-guard-review/simulation-v4/synthetic_slices')
    dataset=analyze(sorted(directory.glob('*.png')),SliceConfig(5,400,400))
    with open('/workspace/ir-image-guard-review/simulation-v4/simulation_300.csv',encoding='utf-8-sig',newline='') as f: source=list(csv.DictReader(f))
    ctl=PowderHoldReference(settings());rows=[];actual=60
    for index in range(1,301):
        feat=dataset.features(index)
        if 84<=index<=90:feat=features(-1)
        if index==60: actual+=2
        layer=index+1 if index>=130 else index
        temp=float(source[index-1]['base_sensor'])
        result=ctl.process(layer,temp,actual,slice_features=feat)
        rows.append((layer,temp,actual,feat,result.diagnostics));actual=result.power_percent
    ns={'m':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'};ET.register_namespace('',ns['m']);ET.register_namespace('r','http://schemas.openxmlformats.org/officeDocument/2006/relationships')
    path=out/'powder_hold_input.xlsx'
    with zipfile.ZipFile(template) as z:
        root=ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
        for row in root.findall('m:sheetData/m:row',ns):
            n=int(row.get('r'));i=n-13
            if not 0<=i<3000:continue
            cells={c.get('r'):c for c in row.findall('m:c',ns)}
            for col in ['A','B','C']+[f'B{chr(ord("L")+j)}' for j in range(12)]:
                addr=f'{col}{n}';cell=cells.get(addr)
                if cell is None:cell=ET.SubElement(row,'{'+ns['m']+'}c',{'r':addr})
                cell.attrib.pop('t',None)
                for child in list(cell):cell.remove(child)
                if i<300:
                    layer,temp,actual,feat,_=rows[i]
                    value={'A':layer,'B':temp,'C':actual}.get(col) if col in ('A','B','C') else feat[FEATURE_NAMES[ord(col[1])-ord('L')]]
                    if value is not None:ET.SubElement(cell,'{'+ns['m']+'}v').text=repr(value)
        xml=ET.tostring(root,encoding='utf-8',xml_declaration=True)
        with zipfile.ZipFile(path,'w') as outzip:
            for item in z.infolist():outzip.writestr(item,xml if item.filename=='xl/worksheets/sheet1.xml' else z.read(item.filename))
    destination=out/'recalculated';destination.mkdir(exist_ok=True)
    profile=out/'lo-profile';soffice=shutil.which('soffice')
    if not soffice:raise RuntimeError('LibreOffice required')
    run=subprocess.run([soffice,'-env:UserInstallation='+profile.resolve().as_uri(),'--headless','--convert-to','xlsx','--outdir',str(destination),str(path)],capture_output=True,text=True,timeout=300)
    if run.returncode:raise RuntimeError(run.stdout+run.stderr)
    book=load_workbook(destination/path.name,read_only=True,data_only=True);sheet=book['수동 테스트']
    legacy=json.loads((ROOT/'app/random80_validation.json').read_text(encoding='utf-8'))['columns'];helpers=('CD','CE','CF','CG','CH','CI','CJ','CK');checks=blank=numeric=0;maxdiff=0;errors=[];future_nonblank=[]
    for cells in sheet.iter_rows():
        rownum=cells[0].row;i=rownum-13
        if 0<=i<300:
            expected=rows[i][4]
            for col in tuple(legacy)+helpers:
                got=cells[column_index_from_string(col)-1].value;want=expected[col];checks+=1
                if want is None:
                    blank+=1
                    if got not in (None,''): errors.append((rownum,col,want,got))
                else:
                    numeric+=1
                    if not isinstance(got,(int,float)):errors.append((rownum,col,want,got));continue
                    diff=abs(got-want);maxdiff=max(maxdiff,diff)
                    if diff>1e-7:errors.append((rownum,col,want,got))
        elif 300<=i<3000:
            for col in ('A','B','C')+tuple('B'+chr(ord('L')+j) for j in range(12)):
                if cells[column_index_from_string(col)-1].value is not None:future_nonblank.append((rownum,col))
    for sh in book:
        for row in sh.iter_rows():
            for cell in row:
                if cell.data_type=='e':errors.append((sh.title,cell.coordinate,'error',cell.value))
    book.close()
    archive=zipfile.ZipFile(template);uncompressed=sum(i.file_size for i in archive.infolist());archive.close()
    evidence=dict(checks=checks,numeric=numeric,blank=blank,maxdiff=maxdiff,errors=errors[:30],error_count=len(errors),future_nonblank=future_nonblank[:10],future_nonblank_count=len(future_nonblank),template_sha256=hashlib.sha256(template.read_bytes()).hexdigest(),template_size=template.stat().st_size,template_uncompressed=uncompressed,rows=300)
    (out/'parity.json').write_text(json.dumps(evidence,indent=2));print(json.dumps(evidence,indent=2))
    if errors or future_nonblank or template.stat().st_size>16*1024**2 or uncompressed>128*1024**2:raise AssertionError('parity or limits failure')

if __name__=='__main__' and len(sys.argv)>1 and sys.argv[1]=='--parity':run_parity(sys.argv[2])
