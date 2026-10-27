"""Build the editable image limiter over the preserved slice.2 workbook.

The source must be the unmodified advisory template (SHA pinned below).
No Excel automation or equipment access is performed by this build tool.
"""
import argparse
from copy import copy
import hashlib
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

SOURCE_SHA256 = '7d870e35c7c288573bf5ab279ede6eaab8fe37633cd0afbcc2c389bcdd79c8b3'
CONFIG = "'이미지 억제 설정'!"


def build(source, output):
    source, output = Path(source), Path(output)
    if hashlib.sha256(source.read_bytes()).hexdigest() != SOURCE_SHA256:
        raise ValueError('Expected the preserved, unchanged slice.2 template')
    book = load_workbook(source)
    sheet = book['수동 테스트']
    config = book.create_sheet('이미지 억제 설정')
    rows = [
        ('항목', '오프라인 시험값', '설명'),
        ('이미지 억제 사용 1 / 중지 0', 1, '이미지가 없으면 기존 계산. 이미지 단독으로 출력 변경하지 않음'),
        ('급락 기준 °C', 4, '직전 3개 연속 관찰의 최댓값 대비 하락'),
        ('급락 전 3개 관찰 범위 °C', 1, '계속 냉각 중인 구간의 오인 억제'),
        ('초기 증출력 보류 관찰 수', 2, '급락을 감지한 관찰을 첫 번째로 포함'),
        ('최대 보호 관찰 수', 24, '종료 후 일반 제어 복귀. 같은 미회복 사건으로 기한 연장하지 않음'),
        ('회복 중 1회 최대 증출력 %p', 1, '감출력은 기존 계산 유지. 누적 출력 상한을 만들지 않음'),
        ('회복 판정 목표 편차 °C', 0.5, '실측 온도가 목표 ±이 값 안에 들어와야 함'),
        ('회복 판정 관찰 간 변화 °C', 0.25, '연속 관찰 간 온도 변화의 절댓값'),
        ('회복 판정 연속 관찰 수', 3, '조기 해제와 다음 급락 재감지 허용 조건'),
        ('사건 후 분사 중 학습 제외 (실험) 1/0', 0, '기본 중지. 사건 후 알려진 분사 구간에서는 계수 갱신만 제외. 일반 출력 계산은 계속'),
        ('설정 유효', None, 'TRUE일 때만 계산. 실제 장비 적용 및 최종 승인은 보류'),
    ]
    for values in rows:
        config.append(values)
    recovery_rows = [
        ('절대 기준 회복 사용 1 / 제한식 비교 0', 1, '급락 전 출력 + 예측 잔여오차 + 제한된 잔여 보정. 현재 출력에 반복 가산하지 않음'),
        ('회복 예측 층수', 8, '양의 온도 회복 추세를 미리 반영해 감속'),
        ('회복 중 1회 최대 감출력 %p', 2, '기존 계산이 더 큰 감출력을 제안하면 기존 감출력 유지'),
        ('회복 이득 상한 %p/°C', 2, '진입 당시 베타/G와 상한 중 작은 값'),
        ('잔여 보정 시작 관찰', 12, '3개 관찰 모두 목표보다 낮고 예측 오차가 남을 때만 양의 보정'),
        ('잔여 보정 시간척도 층', 30, '잔여 보정은 0과 IR 여유 범위 사이. 최종 출력 제한 중 양의 누적 중지'),
        ('최대 절대 회복 관찰', 120, '같은 사건에서 연장하지 않음'),
        ('기존 제어로 복귀 관찰', 12, '회복식과 기존 권장값을 섞고 출력 증감 한도를 점차 복귀'),
        ('실측 안정 복귀 관찰 수', 5, '목표 근처 실측 온도·작은 추세·저장 IR 변화를 함께 확인'),
        ('안정 판정 저장 IR 변화 %p', 1, '회복 예측의 잡음으로 생기는 작은 출력 조정을 허용하는 시험값'),
    ]
    for values in recovery_rows:
        config.append(values)
    config['B12'] = '=AND(COUNT(B2:B11)=10,OR(B2=0,B2=1),B3>0,B4>=0,B5>=1,B5=INT(B5),B6>=B5,B6=INT(B6),B7>=0,B7<=5,B8>=0,B9>=0,B10>=1,B10=INT(B10),OR(B11=0,B11=1),COUNT(B13:B22)=10,OR(B13=0,B13=1),B14>0,B15>0,B15<=5,B16>0,B17>=1,B17=INT(B17),B18>0,B19>=B5,B19>=B17,B19=INT(B19),B20>=1,B20=INT(B20),B21>=1,B21=INT(B21),B22>=0)'
    for col, width in [('A', 37), ('B', 23), ('C', 85)]:
        config.column_dimensions[col].width = width
    old_check = sheet['K5'].value[1:]
    sheet['K5'] = f'=IF({CONFIG}$B$12,{old_check},"이미지 억제 설정 확인")'
    names = {
        'CD': '이미지 억제 문맥 1/0', 'CE': '현재 전체 계획 -1/0/1',
        'CF': '이전 전체 계획 -1/0/1', 'CG': '안정 후 급락 1/0',
        'CH': '이미지 일치 급락 시작', 'CI': '보호 관찰 번호',
        'CJ': '새 급락 감지 허용', 'CK': '실측 회복 연속 관찰',
        'CL': '보호 상태 0일반/1보류/2제한', 'CM': '보호 시간 초과 1/0',
        'CN': '현재 학습 격리', 'CO': '격리 후 깨끗한 관찰 수',
        'CP': 'G·알파·베타 학습 차단', 'CQ': '제한 전 IR 권장값 %',
        'CR': '이 관찰의 목표 °C', 'CS': '회복 연속 중간 계산',
        'CT': '이미지 입력 유효',
        'CU': '급락 기준 이전 온도 °C', 'CV': '측정 하락 폭 °C',
        'CW': '직전 3개 온도 범위 °C', 'CX': '감지·제외 사유 코드',
        'CY': '급락 사건 이력 1/0', 'CZ': '사건 후 분사 학습 제외',
        'DA': '감지용 연속 온도 수', 'DB': '감지 이력 초기화',
        'DC': '급락 전 고정 기준 IR %', 'DD': '진입 당시 회복 이득 %p/°C',
        'DE': '현재 포함 3개 온도 중앙값', 'DF': '2층 온도 추세 °C/층',
        'DG': '회복 온도 오차 °C', 'DH': '예측 잔여오차 °C',
        'DI': '이번 계산 전 잔여 보정 %p', 'DJ': '다음 관찰 잔여 보정 %p',
        'DK': '절대 회복식 제한 전 IR %', 'DL': '기존 계산 복귀 비율 0~1',
        'DM': '회당 양의 출력 한도 %p', 'DN': '증감·고온 반영 회복 IR %',
        'DO': '복귀 혼합 후 제한 전 IR %', 'DP': '안정 복귀 연속 관찰 수',
        'DQ': '이번 관찰 실측 안정 1/0', 'DR': '양의 잔여 보정 허용 1/0',
        'DS': '안정 연속 중간 계산',
    }
    for col, label in names.items():
        sheet[f'{col}12'] = label
        sheet[f'{col}12'].font = Font(bold=True, color='FFFFFF')
        sheet[f'{col}12'].fill = PatternFill('solid', fgColor='365C7C')
        sheet[f'{col}12'].alignment = Alignment(wrap_text=True)
        sheet.column_dimensions[col].width = 22
    for r in range(13, 3013):
        p = r - 1
        valid = f'N{r}="입력 완료"'
        evidence = f'OR(BL{r}=1,BP{r}=1)'
        def formula(col, expression):
            sheet[f'{col}{r}'] = f'=IF({valid},{expression},"")'
        parts = []
        for flag, start, end in [('BL', 'BM', 'BO'), ('BP', 'BQ', 'BS')]:
            parts.append(f'OR(AND(OR(ISBLANK({flag}{r}),{flag}{r}=0),COUNTA({start}{r}:{end}{r})=0),AND(ISNUMBER({flag}{r}),{flag}{r}=1,COUNT({start}{r}:{end}{r})=3,MIN({start}{r}:{end}{r})>=0,MAX({start}{r}:{end}{r})<=100))')
        parts.append(f'OR(ISBLANK(BT{r}),AND(ISNUMBER(BT{r}),BT{r}>=0,BT{r}=INT(BT{r})))')
        for col in ('BU', 'BV', 'BW'):
            parts.append(f'OR(ISBLANK({col}{r}),AND(ISNUMBER({col}{r}),{col}{r}>0))')
        sheet[f'CT{r}'] = '=AND(' + ','.join(parts) + ')'
        old_n = sheet[f'N{r}'].value[1:]
        sheet[f'N{r}'] = f'=IF(B{r}=-1,"온도계 비활성(-1)",IF(NOT(CT{r}),"이미지 입력 확인",{old_n}))'
        if r > 13:
            reset = f'AND({CONFIG}$B$2=1,CY{p}=1,OR({evidence},CD{p}=1),OR(\'수동 테스트\'!$B$5<>CR{p},ABS(C{r}-D{p})>0.000001))'
            formula('O', f'IF(OR(G{r}="새로 시작",N{p}<>"입력 완료",A{r}<>A{p}+1),1,IF({reset},1,MIN(3,O{p}+1)))')
        inherited_context = 'FALSE' if r == 13 else f'AND(N{p}="입력 완료",A{r}=A{p}+1,G{r}<>"새로 시작",CD{p}=1)'
        formula('CD', f'IF({CONFIG}$B$2<>1,0,IF(OR({evidence},{inherited_context}),1,0))')
        changed = 'FALSE' if r == 13 else f'AND(CD{r}=1,OR(\'수동 테스트\'!$B$5<>CR{p},ABS(C{r}-D{p})>0.000001))'
        formula('DB', f'IF(OR(O{r}=1,{changed}),1,0)')
        formula('DA', '1' if r == 13 else f'IF(DB{r}=1,1,MIN(3,DA{p}+1))')
        for col, flag, coverage in [('CE', 'BL', 'BO'), ('CF', 'BP', 'BS')]:
            formula(col, f'IF(OR(CD{r}<>1,{flag}{r}<>1),-1,IF({coverage}{r}>0,1,0))')
        reference = '""' if r < 16 else f'IF(AND(CD{r}=1,DB{r}=0,DA{p}=3),MAX(B{r-3}:B{p}),"")'
        spread = '""' if r < 16 else f'IF(AND(CD{r}=1,DB{r}=0,DA{p}=3),MAX(B{r-3}:B{p})-MIN(B{r-3}:B{p}),"")'
        formula('CU', reference)
        formula('CV', f'IF(ISNUMBER(CU{r}),CU{r}-B{r},"")')
        formula('CW', spread)
        formula('CG', f'IF(AND(CD{r}=1,ISNUMBER(CV{r}),CW{r}<={CONFIG}$B$4,CV{r}>={CONFIG}$B$3),1,0)')
        prior_recovery = '0' if r == 13 else f'IF(DB{r}=1,0,CK{p})'
        recovered = 'FALSE' if r == 13 else f'AND(O{r}<>1,ABS(B{r}-\'수동 테스트\'!$B$5)<={CONFIG}$B$8,ABS(B{r}-B{p})<={CONFIG}$B$9)'
        formula('CS', f'IF(AND(CD{r}=1,{recovered}),{prior_recovery}+1,0)')
        prior_age = '0' if r == 13 else f'IF(DB{r}=1,0,CI{p})'
        prior_armed = '1' if r == 13 else f'IF(DB{r}=1,1,CJ{p})'
        stable = f'IF(O{r}<3,0,IF(AND(CD{r}=1,OR({prior_age}>0,{prior_armed}=0),ABS(B{r}-\'수동 테스트\'!$B$5)<={CONFIG}$B$8,ABS(R{r})<={CONFIG}$B$9,ABS(P{r})<={CONFIG}$B$22),1,0))'
        formula('DQ', stable)
        previous_stable = '0' if r == 13 else f'IF(DB{r}=1,0,DP{p})'
        formula('DS', f'IF(AND({CONFIG}$B$13=1,DQ{r}=1),{previous_stable}+1,0)')
        recovered_state = f'IF({CONFIG}$B$13=1,DS{r}>={CONFIG}$B$21,CS{r}>={CONFIG}$B$10)'
        duration = f'IF({CONFIG}$B$13=1,{CONFIG}$B$19+{CONFIG}$B$20,{CONFIG}$B$6)'
        formula('CM', f'IF(AND(CD{r}=1,NOT({recovered_state}),{prior_age}>={duration}),1,0)')
        formula('CH', f'IF(AND(CD{r}=1,OR({recovered_state},{prior_armed}=1),OR({recovered_state},{prior_age}=0),CM{r}=0,CG{r}=1,CF{r}=1),1,0)')
        formula('CI', f'IF(CD{r}<>1,0,IF(CH{r}=1,1,IF(OR({recovered_state},CM{r}=1,{prior_age}=0),0,{prior_age}+1)))')
        formula('CJ', f'IF(CD{r}<>1,1,IF(CH{r}=1,0,IF({recovered_state},1,{prior_armed})))')
        formula('CK', f'IF(CH{r}=1,0,CS{r})')
        formula('CL', f'IF(CI{r}=0,0,IF(CI{r}<={CONFIG}$B$5,1,IF(AND({CONFIG}$B$13=1,CI{r}>{CONFIG}$B$19),3,2)))')
        formula('DP', f'IF(CH{r}=1,0,DS{r})')
        inherited_seen = 'FALSE' if r == 13 else f'AND(N{p}="입력 완료",A{r}=A{p}+1,G{r}<>"새로 시작",CY{p}=1)'
        formula('CY', f'IF(AND(CD{r}=1,OR(CH{r}=1,{inherited_seen})),1,0)')
        formula('CZ', f'IF(AND({CONFIG}$B$11=1,CY{r}=1,OR(CE{r}=1,CF{r}=1)),1,0)')
        availability_changed = 'FALSE' if r == 13 else f'IF(DB{r}=1,FALSE,OR((CE{r}<>-1)<>(CE{p}<>-1),(CF{r}<>-1)<>(CF{p}<>-1)))'
        formula('CN', f'IF(AND(CD{r}=1,CY{r}=1,OR(CL{r}<>0,CZ{r}=1,AND(CE{r}<>-1,CF{r}<>-1,CE{r}<>CF{r}),{availability_changed})),1,0)')
        clean_limit = "'수동 테스트'!$H$5+3"
        prior_clean = clean_limit if r == 13 else f'IF(DB{r}=1,{clean_limit},CO{p})'
        formula('CO', f'IF(CD{r}<>1,{clean_limit},IF(CN{r}=1,0,MIN({clean_limit},{prior_clean}+1)))')
        formula('CP', f'IF(AND(CD{r}=1,CY{r}=1,OR(CN{r}=1,CO{r}<{clean_limit})),1,0)')
        formula('CR', "'수동 테스트'!$B$5")
        formula('CX', f'IF(CH{r}=1,0,IF(CD{r}=0,1,IF(NOT(ISNUMBER(CU{r})),2,IF(CW{r}>{CONFIG}$B$4,3,IF(CV{r}<{CONFIG}$B$3,4,IF(CF{r}<>1,5,IF(CI{r}>0,6,IF(CJ{r}=0,7,IF(CM{r}=1,8,9)))))))))')
        for col in ('AH', 'AM', 'AN', 'AO', 'AP', 'AV'):
            original = sheet[f'{col}{r}'].value[1:]
            sheet[f'{col}{r}'] = f'=IF(CP{r}=1,0,{original})'
        original = sheet[f'X{r}'].value[1:]
        sheet[f'X{r}'] = f'=IF(CP{r}=1,"",{original})'
        original = sheet[f'D{r}'].value[1:]
        sheet[f'CQ{r}'] = '=' + original
        absolute_active = f'AND({CONFIG}$B$13=1,CL{r}<>0)'
        prior_fixed = '""' if r == 13 else f'IF(DB{r}=1,"",DC{p})'
        prior_gain = '""' if r == 13 else f'IF(DB{r}=1,"",DD{p})'
        formula('DC', f'IF({CONFIG}$B$13<>1,"",IF(CH{r}=1,MEDIAN(C{r-3}:C{p}),{prior_fixed}))')
        formula('DD', f'IF({CONFIG}$B$13<>1,"",IF(CH{r}=1,MIN({CONFIG}$B$16,K{p}/L{p}),{prior_gain}))')
        prior_integral = '0' if r == 13 else f'IF(DB{r}=1,0,DJ{p})'
        formula('DI', f'IF(OR({CONFIG}$B$13<>1,CH{r}=1),0,{prior_integral})')
        for col, expression in (
                ('DE', f'MEDIAN(B{r-2}:B{r})'), ('DF', f'(B{r}-B{r-2})/2'),
                ('DG', f'\'수동 테스트\'!$B$5-DE{r}'),
                ('DH', f'DG{r}-{CONFIG}$B$14*MAX(0,DF{r})'),
                ('DK', f'DC{r}+DD{r}*DH{r}+DI{r}')):
            formula(col, f'IF({absolute_active},{expression},"")')
        formula('DL', f'IF(AND({CONFIG}$B$13=1,CL{r}=3),(CI{r}-{CONFIG}$B$19)/{CONFIG}$B$20,0)')
        formula('DM', f'IF({absolute_active},IF(CL{r}=1,0,MIN(\'수동 테스트\'!$D$6,{CONFIG}$B$7+(\'수동 테스트\'!$D$6-{CONFIG}$B$7)*DL{r})),0)')
        formula('DO', f'IF({absolute_active},(1-DL{r})*DK{r}+DL{r}*CQ{r},"")')
        negative_cap = f'MIN(\'수동 테스트\'!$D$6,{CONFIG}$B$15+(\'수동 테스트\'!$D$6-{CONFIG}$B$15)*DL{r})'
        bounded = f'MIN(\'수동 테스트\'!$H$6,C{r}+DM{r},MAX(\'수동 테스트\'!$F$6,C{r}-{negative_cap},DO{r}))'
        formula('DN', f'IF({absolute_active},IF(B{r}>\'수동 테스트\'!$B$5+\'수동 테스트\'!$D$5,MIN(C{r},{bounded}),{bounded}),"")')
        limiter = f'MIN(\'수동 테스트\'!$H$6,MAX(\'수동 테스트\'!$F$6,C{r}+MIN(V{r},IF(CL{r}=1,0,IF(CL{r}=2,{CONFIG}$B$7,\'수동 테스트\'!$D$6)))))'
        formula('D', f'IF({CONFIG}$B$13=1,IF(CL{r}<>0,IF(CQ{r}<C{r},MIN(DN{r},CQ{r}),DN{r}),CQ{r}),{limiter})')
        eligible = f'IF({absolute_active},IF(AND(CL{r}<>3,CI{r}>={CONFIG}$B$17,MAX(B{r-2}:B{r})<\'수동 테스트\'!$B$5-\'수동 테스트\'!$D$5,DH{r}>\'수동 테스트\'!$D$5,D{r}>=DK{r}-0.000000001),1,0),0)'
        formula('DR', eligible)
        integration = f'DI{r}+IF(DR{r}=1,DD{r}/{CONFIG}$B$18*MAX(0,DH{r}-\'수동 테스트\'!$D$5),0)-IF(B{r}>\'수동 테스트\'!$B$5+\'수동 테스트\'!$D$5,DD{r}/{CONFIG}$B$18*MAX(0,B{r}-\'수동 테스트\'!$B$5),0)'
        formula('DJ', f'IF({absolute_active},MIN(\'수동 테스트\'!$H$6-DC{r},MAX(0,{integration})),DI{r})')
        old_reason = sheet[f'F{r}'].value[1:]
        sheet[f'F{r}'] = f'=IF(N{r}<>"입력 완료",{old_reason},IF({absolute_active},IF(CL{r}=1,"급락 초기 보류",IF(CL{r}=3,"기존 제어로 점진 복귀","절대 기준 회복"))&" · 기준 "&TEXT(DC{r},"0.00")&" → 권장 "&TEXT(D{r},"0.00")&"%",IF(CM{r}=1,"회복 기한 종료 / ",IF(AND(CD{r}=1,OR(CE{r}=-1,CF{r}=-1)),"이미지 미확인 / ",""))&{old_reason}))'
        for col in names:
            sheet[f'{col}{r}'].fill = PatternFill('solid', fgColor='EDF5EB')
    guide = book['이미지 입력 사용법']
    guide['A3'] = 'D열은 실제 회복 권장값입니다. 이미지·급락이 일치하면 급락 전 고정 IR과 진입 당시 베타/G를 사용합니다. 기존 계산 비교값은 CQ열입니다.'
    guide['A5'] = 'BX:CC는 참고 판독, CD:DS는 실제 제어 상태·학습 격리·회복식 근거입니다. 모든 수식과 이미지 억제 설정은 편집할 수 있습니다.'
    guide.append(['이전 전체 이미지에 분사가 있고, 직전 3개 온도가 안정된 뒤 급락할 때만 시작합니다. 중앙 무도포라도 외곽 분사는 교차검증에 포함합니다.'])
    guide.append(['기본 절대 회복식은 급락 전 고정 IR + 진입 베타/G×예측 오차 + 제한된 잔여 보정입니다. 현재 IR에 계속 더하지 않습니다. 초기 2개 관찰은 증출력 보류, 이후 회당 최대 +1/−2%p이며 기존의 더 큰 감출력은 유지합니다.'])
    guide.append(['0~100%와 설정된 최대 증감 한도를 유지합니다. 실측 온도·추세·저장 IR 변화가 5개 연속 안정되면 복귀하며, 최대 120개 회복 관찰 뒤 12개 관찰에 걸쳐 기존 계산으로 점진 복귀합니다. 같은 사건으로 기한을 연장하지 않습니다.'])
    guide.append(['이미지 누락은 미확인으로 남고 진행 중인 보호를 취소하지 않습니다. 무분사 계획만으로 실제 건조 완료나 회복을 판정하지 않습니다.'])
    guide.append(['보호·이미지 계획 전환·이미지 유효성 전환에서는 학습을 격리하며 대기층+3개 관찰 후 재개합니다. 이미지가 계속 미확인이면 보호 종료와 유한 격리 기간 뒤 온도만으로 기존 학습을 재개합니다.'])
    guide.append(['첫 교차검증 급락 이전에는 이미지 계획이나 유효성만으로 기존 출력·학습·수동 변경·목표 변경 동작을 바꾸지 않습니다. 급락을 감지한 관찰부터 학습 격리를 시작합니다.'])
    guide.append(['분사 중 학습 제외는 기본 중지인 실험 옵션입니다. 켜면 실제 감지 사건 이후 알려진 분사 구간에서 G·알파·베타를 유지하므로, 장시간 분사 중 실제 이득 변화에 적응하지 못할 수 있습니다.'])
    guide.append(['잔여 보정 I는 냉각이 층마다 누적된다는 물리 가정이 아닙니다. 예측 후에도 낮은 온도가 이어질 때만 천천히 보정하며, 출력 제한 중 양의 누적을 멈추고 실제 고온에서는 0을 향해 줄입니다. 계산에는 이전 I를 쓰고 갱신 I는 다음 관찰에서 사용합니다.'])
    guide.append(['CU: 기준 온도, CV: 하락 폭, CW: 직전 온도 범위. CX=0감지/1이미지 문맥 없음/2기록 부족/3불안정/4하락 기준 미달/5이전 분사 미확인/6보호 중/7미회복 잠금/8기한 종료/9기타입니다.'])
    guide.append(['목표 변경, 수동 출력 변경, 층 번호 누락·역행, 새로 시작은 이미지 제어 이력을 초기화합니다. 실제 Excel 연결에서는 이전 행 전체를 계산값으로 고정합니다.'])
    guide.append(['이 값들은 오프라인 시험용 잠정값입니다. Excel COM 연동·실제 센서/PLC/램프·장시간 운전은 미검증이며 최종 승인과 장비 적용은 보류입니다.'])
    for row in guide:
        row[0].alignment = Alignment(wrap_text=True, vertical='top')
        guide.row_dimensions[row[0].row].height = 44
    book.calculation.fullCalcOnLoad = True
    book.save(output)
    return output


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('source')
    parser.add_argument('output')
    args = parser.parse_args()
    print(build(args.source, args.output))
