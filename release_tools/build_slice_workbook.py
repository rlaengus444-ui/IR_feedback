"""Build editable advisory workbook from the user's original G/alpha/beta sheet."""
from pathlib import Path
import argparse
from copy import copy
from openpyxl import load_workbook
from openpyxl.formula.translate import Translator
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.utils import get_column_letter

HEADERS = ('현재 이미지 유효(1/0)', '현재 중앙 도포율 %', '현재 중앙 평균강도 %', '현재 전체 도포율 %',
           '이전 이미지 유효(1/0)', '이전 중앙 도포율 %', '이전 중앙 평균강도 %', '이전 전체 도포율 %',
           '이전까지 전체 무분사 연속층', '중앙 반경 mm', 'X DPI', 'Y DPI')

def build(source, output, rows=3000):
    book = load_workbook(source)
    sheet = book['수동 테스트']
    # Retain the original, editable control formulas; remove the old fixed-80 plot objects.
    sheet._charts = []
    for row in range(13, rows + 13):
        if row > 1012:
            for col in range(1, 64):
                old, cell = sheet.cell(1012, col), sheet.cell(row, col)
                cell._style = copy(old._style)
                if isinstance(old.value, str) and old.value.startswith('='):
                    cell.value = Translator(old.value, origin=old.coordinate).translate_formula(cell.coordinate)
        sheet.cell(row, 1, row - 12)
        for col in (2, 3, 7, 8):
            sheet.cell(row, col).value = None
        for col in range(64, 76):
            cell = sheet.cell(row, col)
            cell.value = None
            cell.fill = PatternFill('solid', fgColor='E6F1FF')
            cell.number_format = '0.000'
        # Advisory expressions intentionally do not replace the original D column.
        sheet.cell(row, 76, f'=IF(BP{row}<>1,-1,IF(BS{row}=0,0,IF(BQ{row}=0,1,2)))')
        sheet.cell(row, 77, f'=IF(BX{row}=-1,"이전층 이미지 미확인",IF(BX{row}=0,"계획상 전체 무분사",IF(BX{row}=1,"중앙 미도포·외곽 도포","중앙 도포 참고")))')
        if row >= 16:
            sheet.cell(row, 78, f'=IF(AND(ISNUMBER(B{row}),COUNT(B{row-3}:B{row-1})=3),IF(MAX(B{row-3}:B{row-1})-B{row}>=\'이미지 보조 설정\'!$B$2,1,0),0)')
        else:
            sheet.cell(row, 78, '=0')
        sheet.cell(row, 79, f'=IF(AND(BZ{row}=1,BP{row}=1,BS{row}>0),1,0)')
        sheet.cell(row, 80, f'=IF(AND(BP{row}=1,BS{row}=0,BT{row}>=\'이미지 보조 설정\'!$B$3),1,0)')
        sheet.cell(row, 81, f'=IF(AND(BZ{row}=1,BP{row}=1,BS{row}=0),1,0)')
        for col in range(76, 82):
            sheet.cell(row, col).fill = PatternFill('solid', fgColor='EDF5EB')
    for col, text in enumerate(HEADERS, 64):
        sheet.cell(12, col, text)
    for col, text in enumerate(('이전층 참고 코드(-1/0/1/2)', '이전층 이미지 참고', '급락 참고(1/0)',
                                '급락·이전 도포 계획 일치', '계획상 무분사 연속 참고', '급락·무분사 계획 불일치 참고'), 76):
        sheet.cell(12, col, text)
    for col in range(64, 82):
        sheet.column_dimensions[get_column_letter(col)].width = 18 if col != 77 else 28
        sheet.cell(12, col).alignment = Alignment(wrap_text=True)
        sheet.cell(12, col).font = Font(bold=True, color='FFFFFF')
        sheet.cell(12, col).fill = PatternFill('solid', fgColor='365C7C')
    config = book.create_sheet('이미지 보조 설정')
    for row in [('항목', '시험용 참고값', '의미'), ('급락 참고 기준 °C', 4, '제어 변경값이 아닌 참고 판독 기준'),
                ('무분사 연속 참고층', 3, '실제 건조 완료/분사 종료를 보장하지 않음')]:
        config.append(row)
    config.column_dimensions['A'].width=30;config.column_dimensions['B'].width=20;config.column_dimensions['C'].width=60
    guide = book.create_sheet('이미지 입력 사용법')
    lines = [
        '슬라이싱 데이터 연동 수동 계산 시트',
        '제어 시트: 수동 테스트 / 온도 B13 / 현재 IR C13 / 결과 D13 / 목표 B5 / 층 A13 / 이미지 시작 BL13',
        '기존 G·알파·베타와 권장출력 D열의 수식은 원본 그대로입니다. 새 보조 열은 사용자가 수식 수정 시 참조할 수 있습니다.',
        'BL:BW 파란 열은 프로그램이 현재/이전 이미지 수치를 입력합니다. 운전 중 읽기 전용 데이터로 취급하세요.',
        'BX:CC 녹색 열은 이미지 참고 판독 예시입니다. 도포 계획·중앙 위치와 센서 위치가 완전히 같다고 가정하지 않습니다.',
        '도포율은 유효 분사 픽셀 비율이며, 평균강도는 선택한 색상 기준 0~100 정규화 값입니다. 실제 분사량/실측 건조도는 아닙니다.',
        'T_recoating은 현재 층 분사 전 값입니다. 이전 층의 BP:BS를 우선 대조하고, 현재 BL:BO는 이번 층 분사 계획으로 분리합니다.',
        '유효=0인 이미지는 미확인입니다. 비율의 빈칸을 무분사 0으로 바꾸지 마세요.',
        '파일명 숫자 배정은 누락 번호를 보존합니다. 첫 파일 대응층·DPI·반경·흰색/검정색 분사·기준값을 실제 작업과 확인하세요.',
        '이전 무분사 연속층도 이미지 계획 참고값입니다. 실제 장비의 바인더 분사 실행·잔류 용매 완전 건조를 보장하지 않습니다.',
        '프로그램 연결 중 선택한 원본 파일은 저장하지 않으며, 전용 Excel 복사본에 새 관찰을 순서대로 넣습니다. 첫 관찰 층이 A13에 들어갑니다.',
        '프로그램에서 분석값 포함 시트를 저장하면 오프라인 수동 입력용으로 BL:BW를 미리 채웁니다. 자동 연결에서는 이 입력을 지우고 현재 분석 데이터로 다시 씁니다.',
        '알고리즘을 바꾸려면 D열과 보조 수식을 편집할 수 있습니다. 먼저 계산 미리보기로 확인하세요. 실장비 IR 적용/열 안전성은 검증하지 않았습니다.',
        '층별 수식은 3000개가 준비돼 있습니다. 시작층이 1이 아니어도 A열 실제 층 번호와 이미지 데이터를 맞춰 사용합니다.'
    ]
    for line in lines:guide.append([line])
    guide.column_dimensions['A'].width=130
    for row in guide:
        row[0].alignment=Alignment(wrap_text=True,vertical='top');guide.row_dimensions[row[0].row].height=44
    book.calculation.fullCalcOnLoad=True
    book.save(output)

if __name__=='__main__':
    parser=argparse.ArgumentParser();parser.add_argument('source');parser.add_argument('output');args=parser.parse_args();build(args.source,args.output)
