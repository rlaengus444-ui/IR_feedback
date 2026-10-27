"""Compare actual Python calculations with recalculated XLSX cells offline.

Requires LibreOffice and openpyxl for this development check. This does not
exercise Microsoft Excel COM or any equipment. Output must be a new directory.
"""
from pathlib import Path
import argparse
import csv
import hashlib
import json
import re
import shutil
import subprocess
import sys
import zipfile
from xml.etree import ElementTree as ET

from openpyxl import load_workbook
from openpyxl.utils import column_index_from_string

APP = Path(__file__).resolve().parents[1] / 'app'
sys.path.insert(0, str(APP))
from layer_algorithm import LayerAdaptiveController, LayerSettings
from image_guard import ImageGuardSettings
from dataclasses import replace
from slice_data import SliceConfig, analyze
from slice_workbook import TEMPLATE, save_workbook

NS = {'m': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
ET.register_namespace('', NS['m'])
ET.register_namespace('r', 'http://schemas.openxmlformats.org/officeDocument/2006/relationships')


def patch_inputs(path, rows, clear_future=True):
    with zipfile.ZipFile(path) as archive:
        root = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
        for row in root.findall('m:sheetData/m:row', NS):
            n = int(row.get('r'))
            if not 13 <= n <= 3012:
                continue
            cells = {c.get('r'): c for c in row.findall('m:c', NS)}
            current = rows[n - 13] if n - 13 < len(rows) else None
            if current is None and not clear_future:
                continue
            columns = ['A', 'B', 'C']
            if current is None:
                columns += [f'B{chr(ord("L") + j)}' for j in range(12)]
            for col in columns:
                address = col + str(n)
                cell = cells.get(address)
                if cell is None:
                    cell = ET.SubElement(row, '{' + NS['m'] + '}c', {'r': address})
                cell.attrib.pop('t', None)
                for child in list(cell):
                    cell.remove(child)
                if current is not None:
                    value = {'A': current['layer'], 'B': current['sensor'], 'C': current['previous_ir']}[col]
                    ET.SubElement(cell, '{' + NS['m'] + '}v').text = repr(value)
        xml = ET.tostring(root, encoding='utf-8', xml_declaration=True)
        temp = path.with_suffix('.tmp')
        with zipfile.ZipFile(temp, 'w') as output:
            for item in archive.infolist():
                output.writestr(item, xml if item.filename == 'xl/worksheets/sheet1.xml' else archive.read(item.filename))
    temp.replace(path)


def recalculate(paths, out):
    out.mkdir()
    profile = out.parent / ('lo-profile-' + out.name)
    executable = shutil.which('soffice')
    if not executable:
        raise RuntimeError('LibreOffice is required for this offline check')
    completed = subprocess.run([executable, '-env:UserInstallation=' + profile.as_uri(),
        '--headless', '--convert-to', 'xlsx', '--outdir', str(out), *map(str, paths)],
        text=True, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, timeout=300, check=True)
    print(completed.stdout, flush=True)
    if not all((out / p.name).exists() for p in paths):
        raise RuntimeError('A recalculated workbook is missing')


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--simulation', type=Path, required=True)
    parser.add_argument('--output', type=Path, required=True)
    parser.add_argument('--include-research', action='store_true')
    args = parser.parse_args()
    args.output.mkdir(parents=True, exist_ok=False)
    dataset = analyze(sorted((args.simulation / 'synthetic_slices').glob('*.png')),
                      SliceConfig(radius_mm=5, dpi_x=400, dpi_y=400))
    document = json.loads((APP / 'workbook-replay.settings.json').read_text())
    settings = LayerSettings.from_dict(document)
    guard = ImageGuardSettings.from_dict(document)
    with (args.simulation / 'simulation_300.csv').open(encoding='utf-8-sig', newline='') as stream:
        records = list(csv.DictReader(stream))
    diagnostics = {}
    paths = []
    image_prefix = ('absolute' if 'absolute_sensor' in records[0] else
                    'v3' if 'v3_sensor' in records[0] else 'image')
    modes = [('image', image_prefix, False), ('noimage', 'noimage', False)]
    if args.include_research:
        modes.append(('research', image_prefix + '_print_exclusion_research', True))
    for mode, prefix, research in modes:
        controller = LayerAdaptiveController(settings,
            replace(guard, exclude_print_learning_after_event=True) if research else guard)
        rows, values = [], []
        for record in records:
            row = dict(layer=int(record['layer']), sensor=float(record[prefix + '_sensor']),
                       previous_ir=float(record[prefix + '_previous_ir']))
            result = controller.process(row['layer'], row['sensor'], row['previous_ir'],
                **({'slice_features': dataset.features(row['layer'])} if mode != 'noimage' else {}))
            assert abs(result.power_percent - float(record[prefix + '_ir'])) < 1e-10
            rows.append(row)
            values.append(result.diagnostics)
        diagnostics[mode] = values
        path = args.output / (mode + '.xlsx')
        save_workbook(path, dataset if mode != 'noimage' else None)
        if research:
            workbook = load_workbook(path)
            workbook['이미지 억제 설정']['B11'] = 1
            workbook.save(path)
            workbook.close()
        patch_inputs(path, rows)
        paths.append(path)
    recalculated = args.output / 'recalculated'
    recalculate(paths, recalculated)
    counts = dict(total=0, numeric=0, blank=0, error_cells=0, max_difference=0.)
    for mode, _, _ in modes:
        workbook = load_workbook(recalculated / (mode + '.xlsx'), read_only=True, data_only=True)
        sheet = workbook['수동 테스트']
        for cells in sheet.iter_rows(min_row=13, max_row=312):
            row = cells[0].row
            for col, expected in diagnostics[mode][row - 13].items():
                if not re.fullmatch('[A-Z]{1,2}', col):
                    continue
                actual = cells[column_index_from_string(col) - 1].value
                counts['total'] += 1
                if expected is None:
                    counts['blank'] += 1
                    assert actual in (None, ''), (mode, col, row, expected, actual)
                else:
                    counts['numeric'] += 1
                    assert isinstance(actual, (float, int)), (mode, col, row, expected, actual)
                    difference = abs(actual - expected)
                    counts['max_difference'] = max(counts['max_difference'], difference)
                    assert difference < 1e-7, (mode, col, row, expected, actual, difference)
        for page in workbook:
            for cells in page.iter_rows():
                for cell in cells:
                    if cell.data_type == 'e':
                        counts['error_cells'] += 1
                        raise AssertionError((mode, page.title, cell.coordinate, cell.value))
        workbook.close()
    counts.update(template_sha256=hashlib.sha256(TEMPLATE.read_bytes()).hexdigest(),
                  engine='LibreOffice offline recalculation; Microsoft Excel COM NOT tested',
                  equipment_tested=False)
    (args.output / 'crosscheck.json').write_text(json.dumps(counts, indent=2) + '\n')
    print(json.dumps(counts, indent=2))


if __name__ == '__main__':
    main()
